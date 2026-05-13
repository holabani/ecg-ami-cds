#!/usr/bin/env python3
"""
Generate IEEE-style Major Project viva deck for CardioSense (ecg-ami-cds).
Requires: pip install python-pptx matplotlib numpy
Output: presentations/CardioSense_Major_Project_Viva.pptx
"""

from __future__ import annotations

import io
import os
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
_mpl_cache = ROOT / ".matplotlib-cache"
_mpl_cache.mkdir(exist_ok=True)
os.environ.setdefault("MPLCONFIGDIR", str(_mpl_cache))

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT_DIR = ROOT / "presentations"
OUT_PPTX = OUT_DIR / "CardioSense_Major_Project_Viva.pptx"

NAVY = RGBColor(13, 27, 42)       # #0d1b2a
NAVY_LIGHT = RGBColor(27, 38, 59)
WHITE = RGBColor(255, 255, 255)
SILVER = RGBColor(200, 210, 225)
ACCENT = RGBColor(168, 218, 220)


def synthetic_ecg_like(num_samples: int = 2400) -> tuple[np.ndarray, np.ndarray]:
    t = np.linspace(0, 6, num_samples)
    hr = 1.2
    phase = (t * hr * 2 * np.pi) % (2 * np.pi)
    qrs = np.exp(-((phase - 0.35) ** 2) / 0.002) * 1.4
    qrs += np.exp(-((phase - (2 * np.pi + 0.35)) ** 2) / 0.002) * 1.4
    baseline = 0.08 * np.sin(2 * np.pi * 0.7 * t)
    noise = 0.03 * np.random.randn(num_samples)
    y = baseline + qrs + noise + 0.15 * np.sin(phase)
    return t, y


def fig_to_png_bytes(fig) -> bytes:
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150, bbox_inches="tight", pad_inches=0.05)
    plt.close(fig)
    buf.seek(0)
    return buf.read()


def waveform_background_png() -> bytes:
    np.random.seed(42)
    t, y = synthetic_ecg_like(3200)
    fig = plt.figure(figsize=(14, 8), facecolor="#0d1b2a")
    ax = fig.add_axes([0, 0, 1, 1], facecolor="#0d1b2a")
    ax.plot(t, y, color="#8ebecf", lw=1.2, alpha=0.35)
    ax.plot(t, y * 0.85 + 0.4, color="#6c8ebf", lw=0.9, alpha=0.22)
    ax.set_xlim(t.min(), t.max())
    ax.set_ylim(y.min() - 0.5, y.max() + 0.5)
    ax.axis("off")
    return fig_to_png_bytes(fig)


def pipeline_diagram_png() -> bytes:
    np.random.seed(7)
    t, y = synthetic_ecg_like(1800)
    fig, axes = plt.subplots(2, 1, figsize=(12, 5.2), facecolor="white", height_ratios=[1, 1.1])
    axes[0].plot(t, y, color="#1b263b", lw=1.0)
    axes[0].set_title("Filtered 12-lead (exemplar trace)", fontsize=10, fontname="Times New Roman", color="#1b263b")
    axes[0].set_xticks([])
    axes[0].set_yticks([])
    for spine in axes[0].spines.values():
        spine.set_color("#ccd5e0")

    stages = [
        "Upload / parse\n(CSV · WFDB)",
        "Bandpass 0.5–40 Hz\n+ 50 Hz notch",
        "Pan–Tompkins\n(Lead II, NK2)",
        "Beat alignment\n+ median template",
        "Morph · wavelet\n· inter-lead feats",
        "Torch STFT +\nCNN–BiLSTM",
    ]
    x = np.linspace(0, 1, len(stages))
    axes[1].barh([0] * len(stages), [0.14] * len(stages), left=x - 0.07, height=0.5, color="#415a77", alpha=0.85)
    for xi, lab in zip(x, stages):
        axes[1].text(xi, 0, lab, ha="center", va="center", fontsize=8, fontname="Times New Roman", color="white")
    axes[1].set_xlim(-0.08, 1.08)
    axes[1].set_ylim(-0.55, 0.55)
    axes[1].axis("off")
    axes[1].set_title(
        "Implementation pipeline (backend/preprocessing.py · features.py · model.py)",
        fontsize=10,
        fontname="Times New Roman",
        loc="left",
        color="#1b263b",
    )
    plt.tight_layout()
    return fig_to_png_bytes(fig)


def model_diagram_png() -> bytes:
    fig, ax = plt.subplots(figsize=(11, 4.8), facecolor="white")
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 5)
    ax.axis("off")

    def box(x, y, w, h, text, fc="#415a77"):
        ax.add_patch(
            plt.Rectangle((x, y), w, h, fill=True, facecolor=fc, edgecolor="#1b263b", lw=1.5)
        )
        ax.text(x + w / 2, y + h / 2, text, ha="center", va="center", fontsize=9, color="white", fontname="Times New Roman")

    def arrow(x1, y1, x2, y2):
        ax.annotate("", xy=(x2, y2), xytext=(x1, y1), arrowprops=dict(arrowstyle="-|>", color="#1b263b", lw=1.2))

    box(0.3, 3.5, 1.4, 0.9, "12×T ECG\nmatrix", "#778da9")
    arrow(1.7, 4.0, 2.2, 4.0)
    box(2.2, 3.45, 1.6, 1.0, "STFT\n(n_fft=128\nhop=32)", "#415a77")
    arrow(3.8, 4.0, 4.3, 4.0)
    box(4.3, 3.45, 1.8, 1.0, "CNN branch\n4× Conv2D +\nresidual", "#1b263b")

    arrow(1.7, 3.9, 4.3, 2.9)
    box(2.2, 1.55, 1.6, 1.0, "BiLSTM\n2 layers\n×128 hid", "#415a77")
    arrow(3.8, 2.05, 5.05, 3.45)

    box(6.15, 2.35, 1.35, 1.15, "Attention\nfusion\n(gating)", "#415a77")
    arrow(6.1, 4.0, 6.15, 3.5)
    arrow(6.1, 2.05, 6.15, 2.85)

    arrow(7.5, 2.95, 8.0, 2.95)
    box(8.0, 2.55, 1.35, 0.85, "AMI head\nsigmoid", "#264653")
    box(8.0, 3.45, 1.35, 0.85, "Revasc head\nPCI/CABG", "#264653")

    ax.text(
        5,
        0.55,
        "CardioSense • PyTorch • checkpoint-loaded weights • CPU inference on EC2",
        ha="center",
        fontsize=9,
        fontname="Times New Roman",
        color="#415a77",
    )
    plt.tight_layout()
    return fig_to_png_bytes(fig)


def grafana_style_panel_png() -> bytes:
    fig, axes = plt.subplots(2, 2, figsize=(10.5, 5), facecolor="#121212")
    titles = [
        "ecg_predict_requests_total",
        "ecg_predict_latency_seconds",
        "ecg_ami_risk_prediction_total{tier}",
        "ecg_http_request_duration_seconds",
    ]
    for ax, title in zip(axes.flat, titles):
        ax.set_facecolor("#1e1e1e")
        x = np.arange(12)
        y = np.cumsum(np.random.RandomState(hash(title) % 2**32).randn(12) ** 2 * 0.01 + 0.02)
        ax.fill_between(x, y, alpha=0.35, color="#73BF69")
        ax.plot(x, y, color="#73BF69", lw=1.5)
        ax.set_title(title, color="#e0e0e0", fontsize=8, fontname="Courier New", loc="left")
        ax.tick_params(colors="#aaa", labelsize=7)
        for spine in ax.spines.values():
            spine.set_color("#444")
    fig.suptitle(
        'Illustrative Grafana panels — Prometheus series from backend/metrics.py · dashboard "CardioSense — Overview"',
        color="#cccccc",
        fontsize=9,
        fontname="Times New Roman",
        y=0.98,
    )
    plt.tight_layout(rect=[0, 0, 1, 0.93])
    return fig_to_png_bytes(fig)


def xai_explainability_png() -> bytes:
    """Synthetic visuals aligned with explainability.py outputs (illustrative, not patient data)."""
    rng = np.random.RandomState(42)
    fig = plt.figure(figsize=(12.5, 5.4), facecolor="white")

    ax0 = fig.add_axes([0.045, 0.15, 0.42, 0.72])
    spec = rng.standard_normal((48, 64)) * 0.08
    spec += np.outer(np.linspace(1.0, 0.3, 48), np.linspace(0.4, 1.0, 64))
    im = ax0.imshow(spec, aspect="auto", cmap="RdBu_r", interpolation="bilinear")
    ax0.set_title(
        "Grad-CAM-style saliency (spectrogram branch)",
        fontsize=11,
        fontname="Times New Roman",
        color="#1b263b",
        pad=10,
    )
    ax0.set_xlabel("STFT time frames", fontsize=9, fontname="Times New Roman")
    ax0.set_ylabel("Frequency bins", fontsize=9, fontname="Times New Roman")
    plt.colorbar(im, ax=ax0, fraction=0.046, pad=0.02)

    ax1 = fig.add_axes([0.52, 0.41, 0.44, 0.46])
    leads = ["V6", "V5", "V4", "V3", "V2", "V1", "aVF", "aVL", "aVR", "III", "II", "I"]
    vals = np.array([0.92, 0.81, 0.74, 0.62, 0.51, 0.44, 0.38, 0.31, 0.22, 0.18, 0.14, 0.09])
    cols = plt.cm.Reds(0.35 + 0.65 * (vals / vals.max()))
    ax1.barh(leads, vals, color=cols, edgecolor="#1b263b", linewidth=0.6)
    ax1.set_title(
        "Lead importance · gradcam_lead_importance JSON",
        fontsize=11,
        fontname="Times New Roman",
        color="#1b263b",
    )
    ax1.set_xlim(0, 1.05)
    ax1.tick_params(axis="both", labelsize=8)

    ax2 = fig.add_axes([0.52, 0.12, 0.44, 0.22])
    feats = ["ST deviation V4", "reciprocal aVL", "QRS morphology", "wavelet E_D3", "T symmetry II"]
    shap_like = np.array([0.41, -0.28, 0.19, 0.15, -0.09])
    fcols = ["#c0392b" if v >= 0 else "#2980b9" for v in shap_like]
    ax2.barh(feats, shap_like, color=fcols, edgecolor="#1b263b", linewidth=0.5)
    ax2.axvline(0, color="#555", lw=0.8)
    ax2.set_title(
        "SHAP-style signed features · compute_shap_features()",
        fontsize=10,
        fontname="Times New Roman",
        color="#1b263b",
    )
    ax2.tick_params(axis="both", labelsize=8)

    fig.text(
        0.045,
        0.06,
        "Captum LayerGradCam when installed · heuristic fallback otherwise · narratives via GET /explain/{patient_id}",
        fontsize=9,
        fontname="Times New Roman",
        color="#415a77",
    )

    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=160, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    buf.seek(0)
    return buf.read()


def aws_deploy_diagram_png() -> bytes:
    fig, ax = plt.subplots(figsize=(11, 5), facecolor="white")
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6)
    ax.axis("off")

    def dc(x, y, w, h, title, lines):
        ax.add_patch(plt.Rectangle((x, y), w, h, fill=False, edgecolor="#1b263b", lw=2, linestyle="--"))
        ax.text(x + w / 2, y + h - 0.22, title, ha="center", fontsize=10, fontweight="bold", fontname="Times New Roman")
        yy = y + h - 0.55
        for line in lines:
            ax.text(x + 0.15, yy, line, fontsize=8.5, fontname="Times New Roman", va="top")
            yy -= 0.28

    dc(
        0.35,
        1.2,
        4.2,
        4.5,
        "AWS EC2 host",
        [
            "Docker Compose production stack",
            "docker-compose.production.yml",
            "nginx reverse proxy (TLS) → loopback",
            "Volumes: SQLite · Prometheus TSDB",
        ],
    )

    inner_x, inner_y = 0.95, 1.85

    def svc(ix, iy, w, h, name):
        ax.add_patch(plt.Rectangle((ix, iy), w, h, facecolor="#415a77", edgecolor="#1b263b"))
        ax.text(ix + w / 2, iy + h / 2, name, ha="center", va="center", fontsize=8, color="white", fontname="Times New Roman")

    svc(inner_x, inner_y + 2.5, 1.55, 0.65, "frontend\nNext.js :3000")
    svc(inner_x + 1.75, inner_y + 2.5, 1.55, 0.65, "backend\nFastAPI :8000")
    svc(inner_x, inner_y + 1.55, 1.55, 0.65, "prometheus\n:9090")
    svc(inner_x + 1.75, inner_y + 1.55, 1.55, 0.65, "grafana\n:3000→3010")

    ax.annotate(
        "",
        xy=(6.8, 3.9),
        xytext=(5.15, 3.9),
        arrowprops=dict(arrowstyle="-|>", color="#1b263b", lw=1.5),
    )
    ax.text(
        5.35,
        4.05,
        "HTTPS\n(api · app · grafana)",
        fontsize=8,
        fontname="Times New Roman",
        color="#1b263b",
    )

    ax.text(
        6.85,
        2.8,
        "deploy/nginx/cardiosense-ec2.conf\nproxies to 127.0.0.1 services",
        fontsize=9,
        fontname="Times New Roman",
        color="#415a77",
        bbox=dict(boxstyle="round", facecolor="#e0e6ed"),
    )

    plt.tight_layout()
    return fig_to_png_bytes(fig)


def apply_background(slide, color: RGBColor) -> None:
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color


def set_font_run(run, size_pt: int, bold: bool = False, color: RGBColor = WHITE) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color


def add_title_only_slide(prs: Presentation, title: str, subtitle_lines: list[str]) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    apply_background(slide, NAVY)
    box = slide.shapes.add_textbox(Inches(0.8), Inches(2.3), Inches(11.5), Inches(2.5))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.alignment = PP_ALIGN.CENTER
    set_font_run(p.runs[0], 28, bold=True)

    sub = slide.shapes.add_textbox(Inches(1), Inches(4.55), Inches(11), Inches(2))
    stf = sub.text_frame
    for i, line in enumerate(subtitle_lines):
        para = stf.paragraphs[0] if i == 0 else stf.add_paragraph()
        para.text = line
        para.alignment = PP_ALIGN.CENTER
        set_font_run(para.runs[0], 14, bold=False, color=SILVER)


def add_section_slide(prs: Presentation, title: str, bullets: list[str]) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    apply_background(slide, NAVY)

    hdr = slide.shapes.add_textbox(Inches(0.55), Inches(0.45), Inches(12), Inches(0.85))
    hp = hdr.text_frame.paragraphs[0]
    hp.text = title
    set_font_run(hp.runs[0], 26, bold=True)

    body = slide.shapes.add_textbox(Inches(0.75), Inches(1.35), Inches(11.8), Inches(5.9))
    tf = body.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.text = b
        para.space_after = Pt(10)
        para.level = 0
        set_font_run(para.runs[0], 15, bold=False, color=SILVER)


def add_table_slide(prs: Presentation, title: str, rows: list[list[str]]) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    apply_background(slide, NAVY)

    hdr = slide.shapes.add_textbox(Inches(0.55), Inches(0.45), Inches(12), Inches(0.85))
    hp = hdr.text_frame.paragraphs[0]
    hp.text = title
    set_font_run(hp.runs[0], 26, bold=True)

    cols = len(rows[0])
    rows_n = len(rows)
    tbl_w = Inches(11.5)
    row_h = Inches(0.42)
    left = Inches(0.9)
    top = Inches(1.35)

    table = slide.shapes.add_table(rows_n, cols, left, top, tbl_w, row_h * rows_n).table
    for ri, row in enumerate(rows):
        for ci, cell_text in enumerate(row):
            cell = table.cell(ri, ci)
            cell.text = cell_text
            for p in cell.text_frame.paragraphs:
                for r in p.runs:
                    set_font_run(r, 13, bold=(ri == 0), color=WHITE if ri == 0 else SILVER)
            cell.fill.solid()
            cell.fill.fore_color.rgb = NAVY_LIGHT if ri == 0 else NAVY


def add_image_slide(prs: Presentation, title: str, png_bytes: bytes, caption: str | None = None) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    apply_background(slide, NAVY)

    hdr = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12), Inches(0.75))
    hp = hdr.text_frame.paragraphs[0]
    hp.text = title
    set_font_run(hp.runs[0], 24, bold=True)

    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
        tmp.write(png_bytes)
        tmp.flush()
        slide.shapes.add_picture(tmp.name, Inches(0.65), Inches(1.05), width=Inches(12.1))

    if caption:
        cap = slide.shapes.add_textbox(Inches(0.65), Inches(6.85), Inches(12), Inches(0.55))
        cp = cap.text_frame.paragraphs[0]
        cp.text = caption
        set_font_run(cp.runs[0], 11, bold=False, color=SILVER)


def draw_architecture_slide(prs: Presentation) -> None:
    blank = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank)
    apply_background(slide, NAVY)

    hdr = slide.shapes.add_textbox(Inches(0.55), Inches(0.35), Inches(12), Inches(0.75))
    hp = hdr.text_frame.paragraphs[0]
    hp.text = "System architecture (repository-aligned)"
    set_font_run(hp.runs[0], 24, bold=True)

    legend = slide.shapes.add_textbox(Inches(0.65), Inches(1.05), Inches(12), Inches(1.05))
    lf = legend.text_frame
    lines = [
        "Flow: Clinician → Next.js 14 (TypeScript, Tailwind, Axios) ↔ nginx TLS → FastAPI + Uvicorn (JWT Bearer on protected routes)",
        "Inference path: preprocess → ECGFeatureExtractor → torch CNN-BiLSTM → Grad-CAM + SHAP-style features → SQLAlchemy → SQLite",
        "Observability: Prometheus scrapes GET /metrics · Grafana datasource http://prometheus:9090 · compose stack in docker-compose.production.yml",
    ]
    for i, line in enumerate(lines):
        para = lf.paragraphs[0] if i == 0 else lf.add_paragraph()
        para.text = line
        para.space_after = Pt(6)
        set_font_run(para.runs[0], 12.5, bold=False, color=SILVER)

    # Boxes (positions in inches)
    boxes = [
        (1.0, 2.55, 1.55, 0.62, "User\n(browser)"),
        (3.05, 2.55, 1.85, 0.62, "Next.js 14\nfrontend"),
        (5.35, 2.55, 1.55, 0.62, "nginx\nreverse proxy"),
        (7.35, 2.55, 1.85, 0.62, "FastAPI\n(Uvicorn)"),
        (9.65, 2.55, 1.55, 0.62, "JWT auth\nOTP email"),
        (1.75, 3.95, 2.05, 0.62, "Preproc +\nPan–Tompkins"),
        (4.35, 3.95, 1.95, 0.62, "Features +\nTorch model"),
        (6.75, 3.95, 1.85, 0.62, "XAI\nGrad-CAM · SHAP"),
        (9.05, 3.95, 1.55, 0.62, "SQLite +\nSQLAlchemy"),
        (3.55, 5.35, 2.55, 0.62, "Prometheus → Grafana\n(localhost / tunnel)"),
        (6.85, 5.35, 2.05, 0.62, "AWS EC2 +\nDocker Compose"),
    ]

    for left, top, w, h, label in boxes:
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(left), Inches(top), Inches(w), Inches(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = NAVY_LIGHT
        shape.line.color.rgb = ACCENT
        shape.line.width = Pt(1)
        tf = shape.text_frame
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.text = label
        p.alignment = PP_ALIGN.CENTER
        set_font_run(p.runs[0], 10.5, bold=False, color=WHITE)

    # Simple arrows between top row (coordinates approximate)
    def connector(x1, y1, x2, y2):
        line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        line.line.color.rgb = SILVER
        line.line.width = Pt(1.25)

    connector(2.58, 2.86, 3.02, 2.86)
    connector(4.93, 2.86, 5.33, 2.86)
    connector(6.93, 2.86, 7.32, 2.86)
    connector(9.23, 2.86, 9.62, 2.86)


def main() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)

    prs = Presentation()
    prs.slide_width = Inches(13.333333)
    prs.slide_height = Inches(7.5)

    # ── Slide 1 Cover ──────────────────────────────────────────────
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank)
    apply_background(s1, NAVY)
    bg_png = waveform_background_png()
    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
        tmp.write(bg_png)
        tmp.flush()
        s1.shapes.add_picture(tmp.name, 0, 0, width=prs.slide_width)

    overlay = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, prs.slide_height)
    overlay.fill.solid()
    overlay.fill.fore_color.rgb = NAVY
    overlay.fill.transparency = 0.52
    overlay.line.fill.background()

    tb = s1.shapes.add_textbox(Inches(0.9), Inches(1.35), Inches(11.5), Inches(2.4))
    tfp = tb.text_frame
    p0 = tfp.paragraphs[0]
    p0.text = (
        "AI-Based ECG Analysis for Early Detection of\n"
        "Acute Myocardial Infarction (AMI)\nand Revascularization Need"
    )
    p0.alignment = PP_ALIGN.CENTER
    set_font_run(p0.runs[0], 26, bold=True)

    sub = s1.shapes.add_textbox(Inches(1), Inches(4.05), Inches(11), Inches(2.9))
    sf = sub.text_frame
    meta = [
        "Major Project · CardioSense (ecg-ami-cds codebase)",
        "Bhavya Saxena · Jiya Bajaj · Harmeet Kaur",
        "Faculty Guide: Dr. Renuka Nagpal",
        "Amity School of Engineering and Technology",
        "Amity University",
    ]
    for i, line in enumerate(meta):
        para = sf.paragraphs[0] if i == 0 else sf.add_paragraph()
        para.text = line
        para.alignment = PP_ALIGN.CENTER
        para.space_before = Pt(10) if i == 1 else Pt(0)
        set_font_run(para.runs[0], 14 if i >= 1 else 13, bold=(i == 0), color=SILVER if i > 0 else ACCENT)

    # ── Slide 2 ────────────────────────────────────────────────────
    add_section_slide(
        prs,
        "Problem statement",
        [
            "AMI care quality depends on rapid recognition; delayed diagnosis narrows revascularization benefit.",
            "Standard workflows rely heavily on manual ECG interpretation — expertise-dependent and variable under load.",
            "Interpretation fatigue and communication gaps increase missed STEMI-equivalent patterns.",
            "Hospital-grade CDS rarely bundles deep-learning inference with observability at internship-ready deployment depth.",
            "Clinical stakeholders require transparency (why this risk?) alongside numeric scores.",
        ],
    )

    # ── Slide 3 ────────────────────────────────────────────────────
    add_section_slide(
        prs,
        "Project objectives",
        [
            "Detect AMI likelihood from 12-lead ECG signals via multi-branch CNN–BiLSTM (PyTorch).",
            "Co-predict PCI/CABG (revascularization) need using the shared fused representation.",
            "Expose explainability: Grad-CAM-style lead saliency + SHAP-like morphological/spectral attributions.",
            "Deliver clinician-facing workflows: upload/analyse, history, explain route, alerts UI.",
            "Deploy on AWS EC2 using Docker Compose + nginx with JWT auth and OTP registration.",
            "Instrument Prometheus (/metrics) and Grafana dashboards for latency, traffic, and AMI tier counters.",
        ],
    )

    # ── Slide 4 ────────────────────────────────────────────────────
    add_table_slide(
        prs,
        "Existing vs proposed system",
        [
            ["Aspect", "Existing / conventional", "Proposed · CardioSense implementation"],
            ["Interpretation", "Manual reading; siloed viewers", "AI-assisted scoring + retained human oversight"],
            ["Throughput", "Queue-limited during surge", "FastAPI batch inference path + structured JSON responses"],
            ["Transparency", "Narrative report only", "Grad-CAM lead map + SHAP-feature bars in Next.js UI"],
            ["Monitoring", "Ad hoc logs", "Prometheus scrape + Grafana \"CardioSense — Overview\""],
            ["Deployment", "Workstation scripts", "EC2 · Docker · Compose · nginx TLS ingress"],
        ],
    )

    # ── Slide 5 ────────────────────────────────────────────────────
    draw_architecture_slide(prs)

    # ── Slide 6 ────────────────────────────────────────────────────
    add_image_slide(
        prs,
        "ECG signal processing pipeline",
        pipeline_diagram_png(),
        "NeuroKit2-backed Pan–Tompkins when available; SciPy Butterworth bandpass + notch per preprocessing.py.",
    )

    # ── Slide 7 ────────────────────────────────────────────────────
    add_image_slide(
        prs,
        "Deep learning model (CNN–BiLSTM)",
        model_diagram_png(),
        "Spectrogram CNN branch + temporal BiLSTM + gated fusion; dual sigmoid heads for AMI and revascularization.",
    )

    # ── Slide 8 ────────────────────────────────────────────────────
    add_image_slide(
        prs,
        "Explainable AI — Grad-CAM · SHAP-style attribution",
        xai_explainability_png(),
        "Illustrative panels mirroring predict/explain payload keys stored in Prediction.shap_json / gradcam_json · Next.js renders bars from live API.",
    )

    # ── Slide 9 ────────────────────────────────────────────────────
    add_section_slide(
        prs,
        "Backend API & authentication",
        [
            "FastAPI application main.py — lifespan hooks initialise SQLite via SQLAlchemy models.",
            "POST /predict · POST /predict/upload — JWT-protected; returns ami_probability, revascularization_probability, feature dicts.",
            "POST /auth/register → email OTP · POST /auth/register/verify · POST /auth/login issuing Bearer JWT.",
            "GET /history · GET /alerts · GET /explain/{patient_id} — authorised reads from persisted Prediction rows.",
            "GET /metrics — Prometheus exposition (Counters/Histograms in metrics.py); middleware skips scrape noise.",
            "GET /health — orchestration health-check used by Docker HEALTHCHECK directive.",
        ],
    )

    # ── Slide 10 ───────────────────────────────────────────────────
    add_image_slide(
        prs,
        "Monitoring & observability",
        grafana_style_panel_png(),
        "Series names illustrated above match prometheus_client instrumentation; dashboards provisioned from deploy/grafana/dashboards/.",
    )

    # ── Slide 11 ────────────────────────────────────────────────────
    add_image_slide(
        prs,
        "AWS cloud deployment",
        aws_deploy_diagram_png(),
        "Production Compose binds backend·frontend·Prometheus·Grafana to 127.0.0.1; TLS termination via host nginx configuration.",
    )

    # ── Slide 12 ────────────────────────────────────────────────────
    add_table_slide(
        prs,
        "Results & outputs (no fabricated cohort metrics)",
        [
            ["Deliverable", "What the deployed stack exposes"],
            ["Classification outputs", "AMI probability · revascularization probability · urgency tiering logic"],
            ["XAI artefacts", "gradcam_lead_importance · shap_features JSON mirrored in SQLite Prediction rows"],
            ["Operational signals", "Predict latency histogram · AMI risk-tier counters · optional TP/TN/FP/FN when ami_ground_truth supplied"],
            ["Formal KPI table", "Populate AUROC / Sens / Spec from thesis validation chapter — not hard-coded in repo"],
            ["UI evidence", "Screenshots from Next.js analyse · explain · alerts routes for viva appendix"],
        ],
    )

    # ── Slide 13 ────────────────────────────────────────────────────
    add_section_slide(
        prs,
        "Challenges faced (engineering)",
        [
            "Docker image size vs EC2 RAM — PyTorch + scientific stack lengthened build cycles and demanded tier sizing discipline.",
            "Binding Compose services to loopback while exposing nginx-only ingress complicated smoke-testing.",
            "NEXT_PUBLIC_API_URL correctness — Axios base URL must match TLS-terminated hostname behind proxy.",
            "Prometheus cardinality control — middleware templates route paths to avoid exploding labels with patient IDs.",
            "Robust preprocessing — graceful degradation when SciPy / NeuroKit2 / PyWavelets absent for lightweight CI.",
            "XAI dependency drift — Captum/SHAP optional imports require heuristic fallbacks without breaking API contracts.",
        ],
    )

    # ── Slide 14 ────────────────────────────────────────────────────
    add_section_slide(
        prs,
        "Conclusion & future scope",
        [
            "CNN–BiLSTM fusion captures spectro-temporal AMI cues while emitting dual clinical endpoints in one forward pass.",
            "Explainability hooks (Grad-CAM + SHAP-proxy features) align model outputs with inspectable physiology.",
            "Cloud-native packaging (Compose + nginx + EC2) demonstrates deployable CDS engineering beyond notebook prototypes.",
            "Prometheus/Grafana closes the loop on latency drift and alert storms.",
            "Future scope: wearable streams · mobile clients · multi-site rollout · PostgreSQL migration · transformer encoders.",
            "Thank you.",
        ],
    )

    prs.save(str(OUT_PPTX))
    print(f"Wrote {OUT_PPTX}")


if __name__ == "__main__":
    main()
